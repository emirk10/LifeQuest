import collections
import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    
    # Set to widescreen 16:9
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Common layout styles
    font_title = "Arial"
    font_body = "Arial"
    
    color_title = RGBColor(33, 37, 41) # Dark charcoal
    color_body = RGBColor(73, 80, 87)   # Muted dark gray
    color_placeholder_bg = RGBColor(241, 243, 245) # Light gray
    color_placeholder_border = RGBColor(206, 212, 218) # Border gray
    color_placeholder_text = RGBColor(108, 117, 125) # Gray text
    
    blank_layout = prs.slide_layouts[6]
    
    # Helper to format text box margins to 0 for perfect alignment
    def zero_margins(tf):
        tf.margin_left = Inches(0)
        tf.margin_right = Inches(0)
        tf.margin_top = Inches(0)
        tf.margin_bottom = Inches(0)

    # Helper for adding Title
    def add_title(slide, text):
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.733), Inches(1.0))
        tf = title_box.text_frame
        zero_margins(tf)
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = text
        p.font.name = font_title
        p.font.size = Pt(36)
        p.font.bold = True
        p.font.color.rgb = color_title
        return title_box

    # Helper for adding text block
    def add_bullet_points(slide, points, left, top, width, height):
        txBox = slide.shapes.add_textbox(left, top, width, height)
        tf = txBox.text_frame
        zero_margins(tf)
        tf.word_wrap = True
        
        for i, pt in enumerate(points):
            if i == 0:
                p = tf.paragraphs[0]
            else:
                p = tf.add_paragraph()
                
            p.text = pt
            p.font.name = font_body
            p.font.size = Pt(17)
            p.font.color.rgb = color_body
            p.space_after = Pt(12)
            p.level = 0
                
        return txBox

    # Helper for adding image placeholder
    def add_image_placeholder(slide, description, left, top, width, height):
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = color_placeholder_bg
        shape.line.color.rgb = color_placeholder_border
        shape.line.width = Pt(1.5)
        
        tf = shape.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = f"📸 RESİM EKLE\n\n{description}"
        p.alignment = PP_ALIGN.CENTER
        p.font.name = font_body
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = color_placeholder_text

    # ==========================================
    # SLIDE 1: Kapak
    # ==========================================
    slide1 = prs.slides.add_slide(blank_layout)
    
    # Left blue/gray decorative bar
    accent_bar = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(2.2), Inches(0.1), Inches(3.0))
    accent_bar.fill.solid()
    accent_bar.fill.fore_color.rgb = RGBColor(74, 144, 226) # Nice accent blue
    accent_bar.line.fill.background()
    
    # Title & Subtitle block
    kapak_box = slide1.shapes.add_textbox(Inches(1.2), Inches(2.2), Inches(11.0), Inches(3.5))
    tf1 = kapak_box.text_frame
    zero_margins(tf1)
    tf1.word_wrap = True
    
    p_title = tf1.paragraphs[0]
    p_title.text = "LifeQuest"
    p_title.font.name = font_title
    p_title.font.size = Pt(54)
    p_title.font.bold = True
    p_title.font.color.rgb = color_title
    p_title.space_after = Pt(10)
    
    p_subtitle = tf1.add_paragraph()
    p_subtitle.text = "Yapay Zeka Destekli Kişisel Gelişim ve Oyunlaştırma Platformu"
    p_subtitle.font.name = font_title
    p_subtitle.font.size = Pt(24)
    p_subtitle.font.color.rgb = RGBColor(108, 117, 125)
    p_subtitle.space_after = Pt(40)
    
    p_info = tf1.add_paragraph()
    p_info.text = "Hazırlayan: Mert BAYIR\nDanışman: [Varsa Danışman Adı]"
    p_info.font.name = font_body
    p_info.font.size = Pt(18)
    p_info.font.color.rgb = color_body
    p_info.font.bold = True

    # ==========================================
    # SLIDE 2: İçerik
    # ==========================================
    slide2 = prs.slides.add_slide(blank_layout)
    add_title(slide2, "İçerik")
    
    iceri_list = [
        "• Problem ve Motivasyon",
        "• Çözüm: LifeQuest Nedir?",
        "• Sistem Mimarisi ve Teknoloji Yığını",
        "• Oyunlaştırma Özellikleri (Gamification)",
        "• Yapay Zeka Entegrasyonları",
        "• Sonuç ve Değerlendirme",
        "• Soru-Cevap ve Demo"
    ]
    add_bullet_points(
        slide2, 
        iceri_list, 
        left=Inches(0.8), 
        top=Inches(2.0), 
        width=Inches(11.733), 
        height=Inches(4.8)
    )

    # ==========================================
    # SLIDE 3: Problem ve Motivasyon
    # ==========================================
    slide3 = prs.slides.add_slide(blank_layout)
    add_title(slide3, "Problem ve Motivasyon")
    
    problem_list = [
        "• Kişisel gelişim süreçlerinin ve yeni alışkanlıklar kazanmanın genellikle monoton, sıkıcı ve sürdürülebilir olmaktan uzak bir yapıda olması.",
        "• Kullanıcıların hedeflerine ulaşırken anlık geri bildirim veya ödül mekanizmalarının bulunmaması nedeniyle motivasyonlarını ve disiplinlerini hızla kaybetmeleri.",
        "• İnsanların oyunlaştırma mekanizmalarıyla (XP, seviye, rekabet) çok daha kolay motive olabilmesi ve davranışsal alışkanlıkları eğlenerek kazanabilmesi.",
        "• Geleneksel yapılacaklar listesi (to-do list) ve alışkanlık takip uygulamalarının statik yapısı; kişiye özel esneklik ve duygusal duruma göre uyum sağlama yeteneğinden yoksun olması."
    ]
    add_bullet_points(
        slide3, 
        problem_list, 
        left=Inches(0.8), 
        top=Inches(2.0), 
        width=Inches(11.733), 
        height=Inches(4.8)
    )

    # ==========================================
    # SLIDE 4: Çözüm: LifeQuest Nedir? (With Image)
    # ==========================================
    slide4 = prs.slides.add_slide(blank_layout)
    add_title(slide4, "Çözüm: LifeQuest Nedir?")
    
    cozum_list = [
        "• Kullanıcıların kişisel gelişim hedeflerini başarmalarına yardımcı olan, mikroservis mimarisiyle geliştirilmiş yapay zeka destekli modern bir web platformudur.",
        "• Eğitim, Sağlık, Kariyer ve Sosyal olmak üzere 4 temel kategoride kullanıcılara özel günlük görevler ve hedefler sunar.",
        "• Süreci XP (deneyim puanı), seviye (level), rozet ve başarı (achievement) sistemleriyle eğlenceli ve sürdürülebilir bir oyuna dönüştürür.",
        "• Kullanıcıların başarılarını paylaşabildiği, görsel görev etiketleme yapabildiği ve arkadaşlarından Kudos (takdir puanı) kazanabildiği sosyal bir ekosistem sağlar."
    ]
    add_bullet_points(
        slide4, 
        cozum_list, 
        left=Inches(0.8), 
        top=Inches(2.0), 
        width=Inches(6.0), 
        height=Inches(4.8)
    )
    add_image_placeholder(
        slide4, 
        "LifeQuest Ana Sayfa veya Dashboard genel görüntüsü", 
        left=Inches(7.2), 
        top=Inches(2.0), 
        width=Inches(5.333), 
        height=Inches(4.5)
    )

    # ==========================================
    # SLIDE 5: Sistem Mimarisi ve Teknolojiler (With Image)
    # ==========================================
    slide5 = prs.slides.add_slide(blank_layout)
    add_title(slide5, "Sistem Mimarisi ve Teknolojiler")
    
    mimar_list = [
        "• Sistem 4 temel bağımsız bileşenden oluşan konteynerize mikroservis mimarisi ile kurulmuştur (Frontend, Backend, AI Servisi, Veritabanı).",
        "• Docker ve Docker Compose kullanılarak tüm servisler izole edilmiş şekilde orkestre edilir ve kolayca dağıtılır.",
        "• Backend: .NET 9.0, Onion Architecture (Soğan Mimarisi), CQRS (MediatR), Entity Framework Core ORM ve SignalR ile real-time iletişim altyapısı.",
        "• Frontend: React 19 ve hızlı derleme için Vite aracı, tasarımlar için modern ve responsive TailwindCSS kütüphanesi.",
        "• Yapay Zeka Servisi: Python 3.11 ve FastAPI tabanlı, makine öğrenmesi modelleri ve Türkçe NLP boru hattı (Pipeline) barındıran mikroservis."
    ]
    add_bullet_points(
        slide5, 
        mimar_list, 
        left=Inches(0.8), 
        top=Inches(2.0), 
        width=Inches(6.0), 
        height=Inches(4.8)
    )
    add_image_placeholder(
        slide5, 
        "4 temel bileşeni gösteren basit bir mimari şeması", 
        left=Inches(7.2), 
        top=Inches(2.0), 
        width=Inches(5.333), 
        height=Inches(4.5)
    )

    # ==========================================
    # SLIDE 6: Oyunlaştırma (Gamification) Özellikleri (With Image)
    # ==========================================
    slide6 = prs.slides.add_slide(blank_layout)
    add_title(slide6, "Oyunlaştırma (Gamification) Özellikleri")
    
    game_list = [
        "• Görev Tabanlı İlerleme: Görevlerin zorluk seviyelerine göre (Kolay 1.0x, Orta 1.2x, Zor 1.5x) XP (Deneyim Puanı) kazanımı ve SignalR ile anlık seviye takibi.",
        "• Seviye Sistemi: Her 1000 XP barajında kullanıcının genel seviyesinin (General Level) artması ve kategori bazlı seviye ilerleme çubukları.",
        "• Rozet ve Başarı Sistemi: Common, Rare, Epic ve Legendary nadirlik derecelerinde tasarlanmış başarı rozetleri ve profil sayfasında sergilenme alanı.",
        "• Streak (Seri) Sistemi: Kullanıcıların üst üste görev tamamlama sürekliliğini takip eden ve motivasyonu koruyan aktif gün sayacı (Streak) mekanizması."
    ]
    add_bullet_points(
        slide6, 
        game_list, 
        left=Inches(0.8), 
        top=Inches(2.0), 
        width=Inches(6.0), 
        height=Inches(4.8)
    )
    add_image_placeholder(
        slide6, 
        "Uygulamadaki Profil sayfasından kazanılan rozetleri veya XP çubuğunu gösteren ekran görüntüsü", 
        left=Inches(7.2), 
        top=Inches(2.0), 
        width=Inches(5.333), 
        height=Inches(4.5)
    )

    # ==========================================
    # SLIDE 7: Yapay Zeka Entegrasyonu - Görev Önerisi (With Image)
    # ==========================================
    slide7 = prs.slides.add_slide(blank_layout)
    add_title(slide7, "Yapay Zeka Entegrasyonu - Görev Önerisi")
    
    ai_rec_list = [
        "• Kişisel Onboarding: Yeni kullanıcılar kayıt esnasında 7 soruluk \"Seni Tanıyalım\" anketine tabi tutularak kişisel hedef ve kapasite profili oluşturulur.",
        "• Ağırlıklı Öneri Formülü: Hedef Eşleşmesi (%40), Güçlü Yönler (%25), Boş Zaman Müsaitliği (%20) ve Zorluk Tercihi (%15) kriterlerine göre hibrit puanlama yapılır.",
        "• Hibrit Öneri Motoru: Kayıtlı kullanıcılar için Item-Based Collaborative Filtering (İşbirlikçi Filtreleme) kullanılırken, yeni kullanıcılar için İçerik Tabanlı (Content-Based) soğuk başlangıç (Cold Start) fallback algoritması çalıştırılır.",
        "• Günlük Akıllı Atama: Algoritma sonucunda kullanıcının günlük tamamlayabileceği, kişisel gelişimine en uygun 3 akıllı görev her sabah otomatik atanır."
    ]
    add_bullet_points(
        slide7, 
        ai_rec_list, 
        left=Inches(0.8), 
        top=Inches(2.0), 
        width=Inches(6.0), 
        height=Inches(4.8)
    )
    add_image_placeholder(
        slide7, 
        "\"Seni Tanıyalım\" anketi veya AI Görev Önerilerini gösteren kartların resmi", 
        left=Inches(7.2), 
        top=Inches(2.0), 
        width=Inches(5.333), 
        height=Inches(4.5)
    )

    # ==========================================
    # SLIDE 8: Yapay Zeka Entegrasyonu - AI Chatbot (With Image)
    # ==========================================
    slide8 = prs.slides.add_slide(blank_layout)
    add_title(slide8, "Yapay Zeka Entegrasyonu - AI Chatbot")
    
    ai_chat_list = [
        "• Türkçe Doğal Dil İşleme (NLP): Karakter n-gram seviyesinde TF-IDF vektörleştirme ve Türkçe stop-words temizliği yapan özel preprocess (ön işleme) adımları.",
        "• Dinamik Sınıflandırıcılar: Arka planda Logistic Regression tabanlı 19 sınıflı Niyet Sınıflandırıcı (Intent) ve 3 sınıflı Duygu Analizörü (Sentiment) modelleri eşzamanlı çalışır.",
        "• Davranışsal Bağlam Duyarlılığı: Chatbot, kullanıcının o anki duygusal durumuna ve profil verilerine (seviye, streak, zayıf yönler) göre proaktif dürtme (Nudge) mesajları ve yönlendirmeleri yapar.",
        "• Kriz ve Aciliyet Tespiti: Level 3 kriz durumlarında görev önermeyi durdurup ALO 183 Aile ve Sosyal Hizmetler gibi destek hatlarına yönlendirme emniyet filtresi."
    ]
    add_bullet_points(
        slide8, 
        ai_chat_list, 
        left=Inches(0.8), 
        top=Inches(2.0), 
        width=Inches(6.0), 
        height=Inches(4.8)
    )
    add_image_placeholder(
        slide8, 
        "Chatbot ile yapılan bir mesajlaşmanın ekran görüntüsü", 
        left=Inches(7.2), 
        top=Inches(2.0), 
        width=Inches(5.333), 
        height=Inches(4.5)
    )

    # ==========================================
    # SLIDE 9: Sonuç ve Öğrenilenler
    # ==========================================
    slide9 = prs.slides.add_slide(blank_layout)
    add_title(slide9, "Sonuç ve Öğrenilenler")
    
    sonuc_list = [
        "• Başarıyla Geliştirilen Sistem: Modern yazılım mimarileri (Microservices, CQRS, Onion Architecture) ve Docker üzerinde koordine edilen 26 tabloluk, tam kapsamlı ve yapay zeka destekli bir kişisel gelişim platformu başarıyla tamamlandı.",
        "• Teknolojik Entegrasyon Deneyimi: .NET Web API, React UI ve Python FastAPI servislerinin Docker Compose orkestrasyonuyla güvenli haberleşmesi (httpOnly Cookie, Refresh Token) ve NLP modellerinin web entegrasyonu sağlandı.",
        "• Kazanılan Değer: Gelişim ve disiplin süreçleri sıkıcı bir rutin olmaktan çıkarılarak, sürdürülebilir, oyunlaştırılmış ve kullanıcıların birbirini motive ettiği sosyal bir alışkanlığa dönüştürüldü."
    ]
    add_bullet_points(
        slide9, 
        sonuc_list, 
        left=Inches(0.8), 
        top=Inches(2.0), 
        width=Inches(11.733), 
        height=Inches(4.8)
    )

    # ==========================================
    # SLIDE 10: Soru - Cevap ve Demo (With Image)
    # ==========================================
    slide10 = prs.slides.add_slide(blank_layout)
    add_title(slide10, "Soru - Cevap ve Demo")
    
    demo_list = [
        "• Dinlediğiniz için teşekkürler.",
        "• Sorularınız?",
        "• Canlı Demo Kapsamı: Yeni kullanıcı kaydı oluşturma, \"Seni Tanıyalım\" anketi, akıllı günlük görevlerin başlatılması (Timer) ve AI Chatbot ile etkileşim adımları."
    ]
    add_bullet_points(
        slide10, 
        demo_list, 
        left=Inches(0.8), 
        top=Inches(2.0), 
        width=Inches(6.0), 
        height=Inches(4.8)
    )
    add_image_placeholder(
        slide10, 
        "Uygulamanın Karanlık Mod (Dark Mode) görüntüsü veya sade bir Teşekkürler görseli", 
        left=Inches(7.2), 
        top=Inches(2.0), 
        width=Inches(5.333), 
        height=Inches(4.5)
    )

    # Save presentation
    prs.save(r"C:\Users\mthng\OneDrive\Desktop\LifeQuest_Bitirme_Sunumu.pptx")
    print("PowerPoint sunumu basariyla guncellendi!")

if __name__ == '__main__':
    create_presentation()
