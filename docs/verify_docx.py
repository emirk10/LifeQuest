"""Export updated docx to text for verification."""
import os
from docx import Document

DST = os.path.join(os.path.dirname(__file__), "yeni_updated.docx")
OUT = os.path.join(os.path.dirname(__file__), "doc_updated_content.txt")
doc = Document(DST)

with open(OUT, 'w', encoding='utf-8') as f:
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip():
            f.write(f"[{i:4d}] {p.text}\n")

print(f"Exported to {OUT}")
print(f"Total non-empty paragraphs: {sum(1 for p in doc.paragraphs if p.text.strip())}")
