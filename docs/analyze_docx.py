"""Analyze the docx structure: headings, paragraph styles, images, and their positions."""
import sys
import io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

from docx import Document
import os

doc_path = os.path.join(os.path.dirname(__file__), "yeni.docx")
doc = Document(doc_path)

output_path = os.path.join(os.path.dirname(__file__), "doc_structure.txt")

with open(output_path, 'w', encoding='utf-8') as f:
    f.write("=" * 80 + "\n")
    f.write("PARAGRAPH ANALYSIS\n")
    f.write("=" * 80 + "\n")

    for i, para in enumerate(doc.paragraphs):
        style_name = para.style.name if para.style else "None"
        text_preview = para.text[:120].strip() if para.text.strip() else "(empty)"
        
        # Check for inline images in paragraph XML
        from lxml import etree
        ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
              'wp': 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing',
              'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
              'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
        
        drawings = para._element.findall('.//wp:inline', ns) + para._element.findall('.//wp:anchor', ns)
        img_marker = f" [IMAGE x{len(drawings)}]" if drawings else ""
        
        if text_preview != "(empty)" or img_marker:
            f.write(f"  [{i:4d}] Style: {style_name:35s} | {text_preview}{img_marker}\n")

    f.write("\n" + "=" * 80 + "\n")
    f.write("ALL STYLES USED\n")
    f.write("=" * 80 + "\n")
    styles_used = set()
    for para in doc.paragraphs:
        if para.style:
            styles_used.add(para.style.name)
    for s in sorted(styles_used):
        f.write(f"  - {s}\n")

    f.write("\n" + "=" * 80 + "\n")
    f.write("IMAGES/RELATIONSHIPS\n")
    f.write("=" * 80 + "\n")
    for rel in doc.part.rels.values():
        if "image" in rel.reltype:
            f.write(f"  rId: {rel.rId}, Target: {rel.target_ref}\n")

    f.write(f"\nTotal paragraphs: {len(doc.paragraphs)}\n")

print(f"Structure written to {output_path}")
