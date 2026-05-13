"""
Part 1: ÖZET, BÖLÜM 1 (Giriş, 1.3, 1.4) ve BÖLÜM 2.1 (Formülasyon) güncellemeleri.
Geçmiş zaman kullanılacak. Bootstrap/React Native kaldırılacak.
"""
import shutil, os
from docx import Document

SRC = os.path.join(os.path.dirname(__file__), "yeni.docx")
DST = os.path.join(os.path.dirname(__file__), "yeni_updated.docx")

# Yedek al
shutil.copy2(SRC, DST)
doc = Document(DST)

def replace_para_text(para, new_text):
    """Paragrafın metnini değiştir, ilk run'ın formatını koru."""
    if not para.runs:
        para.text = new_text
        return
    fmt = para.runs[0].font
    style_name = para.style.name if para.style else None
    # Tüm run'ları temizle
    for run in para.runs:
        run.text = ""
    para.runs[0].text = new_text

def find_para_by_text(doc, search_text, start_idx=0):
    """Paragrafı metin aramarak bul, index döndür."""
    for i in range(start_idx, len(doc.paragraphs)):
        if search_text in doc.paragraphs[i].text:
            return i
    return -1

# ============================================================
# DÖNEM: "Güz" → "Bahar"
# ============================================================
idx = find_para_by_text(doc, "2025 - 2026")
if idx >= 0:
    replace_para_text(doc.paragraphs[idx], "2025 - 2026 Bahar Dönemi")
    print(f"[{idx}] Dönem güncellendi → Bahar")

# ============================================================
# BSM 401 → BSM 498
# ============================================================
idx = find_para_by_text(doc, "BSM 401")
if idx >= 0:
    replace_para_text(doc.paragraphs[idx], "BSM 498 BİTİRME ÇALIŞMASI")
    print(f"[{idx}] BSM 401 → BSM 498 güncellendi")

# ============================================================
# ÖZET (paragraf 119-127) — Geçmiş zaman, mobil kaldır, teknoloji güncelle
# ============================================================
ozet_texts = {
    119: 'Günümüzün hızlı tempolu dijital dünyası, bireylerin dikkat dağınıklığı yaşamasına ve uzun vadeli kişisel gelişim hedeflerine (eğitim, sağlık, kariyer) odaklanmakta zorlanmasına neden olmaktadır. Bu durum, bireylerin motivasyon sürekliliği konusunda güçlükler yaşamasına ve potansiyellerini tam kullanamamasına yol açmaktadır. Bu proje, "Yaşam Pusulası" adlı bir web platformu aracılığıyla, bu yaygın motivasyon sorununa oyunlaştırma (gamification) dinamiklerini kullanarak yenilikçi bir çözüm sunmayı amaçlamıştır.',

    121: 'Projenin temel amacı, kullanıcıların hedeflerini (eğitim, sağlık, kariyer ve sosyal sorumluluk gibi) görevlere bölmelerini ve tamamladıkça XP/Seviye kazanmalarını sağlamaktır. Uygulamanın özgün yönü, bireysel gelişimi toplumsal fayda ile bütünleştiren ve Öz Belirlenim Teorisi\'ne dayanan kategori bazlı seviyeleme mekanizmasıdır. Bu yaklaşım ile kullanıcı motivasyonunun geleneksel hedef takip uygulamalarına göre daha efektif olarak artırılması; kullanıcıların kişisel gelişiminin ve uygulamayı kullanırken kazanacakları toplumsal bilincin daha etkili bir şekilde artması hedeflenmiştir.',

    123: 'Yöntem olarak, kullanıcının mevcut seviyesine ve ruh haline göre kural tabanlı (rule-based) bir sistem görev ataması yapmıştır. Gelişim süreci, bağımsız bir Python FastAPI mikro servisi olarak geliştirilen yapay zekâ destekli chatbot ile desteklenmiştir. Bu chatbot; niyet sınıflandırma (intent classification), duygu analizi (sentiment analysis) ve aciliyet tespiti (urgency detection) modülleri ile kullanıcıya bağlam-duyarlı motivasyonel geri bildirimler sunmuştur. Sistem, ASP.NET Core API (Clean Architecture) üzerinde çalışmış, PostgreSQL veritabanı kullanmış ve arayüz Vite + React 19 + TailwindCSS ile tek sayfa uygulama (SPA) olarak geliştirilmiştir.',

    125: 'Bu proje, Birleşmiş Milletler Sürdürülebilir Kalkınma Amaçları (SKA) arasında yer alan SKA 3 (Sağlıklı ve Kaliteli Yaşam), SKA 4 (Nitelikli Eğitim) ve SKA 11 (Sürdürülebilir Şehirler ve Topluluklar) hedeflerine katkı sağlamayı hedeflemiştir. Ulusal ölçekte ise projemiz, T.C. On İkinci Kalkınma Planı (2024-2028) ile doğrudan ilişkilidir. Planda yer alan 675. madde ile uyumlu olarak platform, "hayat boyu öğrenme sisteminin işgücü piyasasının talep ettiği becerileri kazandırması ve bu becerilerin iş ve sosyal yaşamda etkin kullanımı" hedefini (projemizin Eğitim ve Kariyere yönelik kategoriler aracılığıyla) desteklemiştir. Aynı zamanda, projemizin yenilikçi toplumsal fayda boyutu, Plandaki 746.11. madde olan "Gençlerin gönüllülük faaliyetlerine ilişkin bilinçlendirilmesi ve katılımlarının artırılmasına yönelik çalışma yürütülecektir" hedefiyle (projemizin Sosyal Sorumluluk kategorisi aracılığıyla) tam olarak örtüşmüştür. Bu yaklaşım, teknolojiyi insan merkezli bir dönüşüm aracı olarak konumlandıran Toplum 5.0 vizyonuna da hizmet etmiştir.',

    127: 'Sonuç olarak, bu uygulama bireysel hedef belirleme sürecini toplumsal bilinçle birleştirerek dijitalleştirmiş; kullanıcıların motivasyon ve görev tamamlama oranlarını artırarak Türkiye\'nin dijital dönüşüm hedeflerine somut bir katkı sunma potansiyeli taşıyan çalışan bir prototip olarak hayata geçirilmiştir.',
}

for idx, text in ozet_texts.items():
    replace_para_text(doc.paragraphs[idx], text)
    print(f"[{idx}] ÖZET paragrafı güncellendi")

# ============================================================
# 1.3 Yöntem ve Seçim Sebepleri (paragraf 162, 164)
# ============================================================
idx = find_para_by_text(doc, "Bu çalışmada yazılım mimarisi olarak")
if idx >= 0:
    replace_para_text(doc.paragraphs[idx], 'Bu çalışmada yazılım mimarisi olarak "Temiz Mimari" (Clean Architecture) yaklaşımı benimsenmiştir. Bu yöntemin seçilme sebebi, iş kurallarını (çekirdek katman) dış bağımlılıklardan (veritabanı, arayüz) izole ederek sistemin test edilebilirliğini ve sürdürülebilirliğini artırmaktır. Karmaşık iş akışlarının yönetimi için CQRS (Command Query Responsibility Segregation) tasarım deseni ve MediatR kütüphanesi kullanılmıştır. Bu desen, okuma ve yazma işlemlerini ayırarak performans optimizasyonu sağlamıştır.')
    print(f"[{idx}] 1.3 ilk paragraf güncellendi")

idx = find_para_by_text(doc, "Veri tabanı yönetiminde PostgreSQL")
if idx >= 0:
    replace_para_text(doc.paragraphs[idx], 'Veri tabanı yönetiminde PostgreSQL, nesne-ilişkisel eşleme aracı olarak Entity Framework Core tercih edilmiştir. İstemci tarafında ise Web arayüzü için Vite + React 19 + TailwindCSS teknolojileri kullanılmıştır. Bu teknolojilerin seçilme nedeni, bileşen tabanlı yapıları sayesinde kod tekrarını azaltmaları ve geniş topluluk desteğine sahip olmalarıdır. Yapay zeka asistanı için Python FastAPI tabanlı bağımsız bir mikro servis geliştirilmiş olup, bu servis scikit-learn kütüphanesi ile eğitilmiş makine öğrenmesi modelleri (SVM) kullanmıştır. Ayrıca, kullanıcıya uygun görevlerin atanmasında "kural tabanlı" (rule-based) bir algoritma kullanılmış ve karar mekanizmalarının esnek ve güncellenebilir olması sağlanmıştır.')
    print(f"[{idx}] 1.3 ikinci paragraf güncellendi")

# ============================================================
# 1.4 Raporun Akışı (paragraf 173 — mobil kaldır)
# ============================================================
idx = find_para_by_text(doc, "Geliştirilen mobil ve web arayüz")
if idx >= 0:
    replace_para_text(doc.paragraphs[idx], 'Bölüm 3 (Tasarım Detayları ve İş Akışları): Geliştirilen web arayüz tasarımları, kodlama süreçleri, sistemin kurulum aşamaları bu bölümde sunulmuştur.')
    print(f"[{idx}] 1.4 Bölüm 3 açıklaması güncellendi")

# ============================================================
# 2.1 Formülasyon (paragraf 205, 207, 211, 219, 224)
# ============================================================
idx = find_para_by_text(doc, "Projenin temel mekanizması olan")
if idx >= 0:
    replace_para_text(doc.paragraphs[idx], 'Projenin temel mekanizması olan "Genel Seviye" (General Level) hesaplaması, kullanıcının toplam deneyim puanına (TotalXP) dayalı kademeli bir fonksiyon ile gerçekleştirilmiştir. Bu yaklaşımda, başlangıç seviyelerinde hızlı ilerleme sağlanırken, ileri seviyelerde daha fazla çaba gerektirilmiştir. Sistemde ayrıca her kategori (C) kullanıcının o alandaki yetkinliğini temsil eden ayrı bir seviye değerine (Lc) sahiptir.')
    print(f"[{idx}] 2.1 Genel seviye açıklaması güncellendi")

idx = find_para_by_text(doc, "Bu işlem aşağıdaki denklem ile ifade")
if idx >= 0:
    replace_para_text(doc.paragraphs[idx], 'Genel seviye hesaplaması aşağıdaki kademeli fonksiyon ile gerçekleştirilmiştir:')
    print(f"[{idx}] 2.1 denklem tanıtım cümlesi güncellendi")

# Formül görselinin olduğu paragraftan sonra açıklama (paragraf 211)
idx = find_para_by_text(doc, "Burada; Lgenel")
if idx >= 0:
    replace_para_text(doc.paragraphs[idx], 'Seviye 1→2: 500 XP, Seviye 2→3: 1000 XP, Seviye 3→5: 2500 XP (1250 XP/seviye), Seviye 5→10: 11000 XP (2200 XP/seviye), Seviye 10+: Her ek seviye için 5000 XP. Bu kademeli artış yapısı, kullanıcının başlangıçta hızlı ilerlemesini sağlarken ileri seviyelerde daha fazla çaba gerektirerek motivasyon sürekliliğini korumuştur.')
    print(f"[{idx}] 2.1 seviye açıklaması güncellendi")

idx = find_para_by_text(doc, "Ayrıca bir üst seviyeye geçmek için gereken")
if idx >= 0:
    replace_para_text(doc.paragraphs[idx], 'Ayrıca kategori bazlı seviye atlamak için gereken deneyim puanı (XpToNextLevel) doğrusal bir artış modeli ile hesaplanmıştır. Her kategori seviyesinde bir üst seviyeye geçmek için gereken eşik, sabit bir miktar (50 XP) kadar artırılmıştır. Başlangıç seviyesinde 100 XP gerekliyken, her sonraki seviyede bu değer 50 XP artarak devam etmiştir (150, 200, 250...).')
    print(f"[{idx}] 2.1 XP artış açıklaması güncellendi")

idx = find_para_by_text(doc, "Burada; XPtaban sabit")
if idx >= 0:
    replace_para_text(doc.paragraphs[idx], 'Ek olarak, görev tamamlama sürecinde XP Telafi Oranı (XpCompensationRate) mekanizması uygulanmıştır. Normal tamamlamada oran 1.0 (tam XP), ertelenen görevlerde oran 0.70 (−30% ceza), erteleme sonrası telafi tamamlamada ise +%10 bonus uygulanmıştır. Sistemde ayrıca ardışık gün görev tamamlama takibi (streak) mekanizması bulunmakta olup, kullanıcının motivasyonunu sürdürmek amacıyla uzun seriler ödüllendirilmiştir. Kullanıcının günlük görev limiti ise seviyesine göre dinamik olarak hesaplanmış olup, başlangıçta 3 görev ile başlayıp seviye ilerledikçe maksimum 10 göreve kadar artırılmıştır.')
    print(f"[{idx}] 2.1 XP telafi açıklaması güncellendi")

# ============================================================
# 2.2 Donanım Mimarisi (paragraf 228, 230, 232-234)
# ============================================================
idx = find_para_by_text(doc, "istemci-sunucu (client-server) mimarisi")
if idx >= 0:
    replace_para_text(doc.paragraphs[idx], '"Yaşam Pusulası" projesi, konteyner tabanlı bir mikro servis mimarisi üzerinde çalışmıştır. Docker Compose ile orkestre edilen beş ayrı servis konuşlandırılmıştır: PostgreSQL veritabanı, ASP.NET Core Backend API, Python FastAPI AI Mikro Servisi, Vite/React Frontend uygulaması ve geliştirme ortamında e-posta testi için MailHog servisi.')
    print(f"[{idx}] 2.2 Donanım mimarisi giriş güncellendi")

idx = find_para_by_text(doc, "Sunucu Tarafı: Uygulamanın arka ucu")
if idx >= 0:
    replace_para_text(doc.paragraphs[idx], 'Sunucu Tarafı: Uygulamanın arka ucu (Backend API), yapay zeka servisi ve veritabanı hizmetleri Docker konteynerleri içerisinde çalıştırılmıştır. PostgreSQL veritabanı kalıcı veri depolama (volume mount) ile yapılandırılmış, bağlantı havuzlama (connection pooling) ve sağlık kontrolü (healthcheck) mekanizmaları devreye alınmıştır.')
    print(f"[{idx}] 2.2 Sunucu tarafı güncellendi")

idx = find_para_by_text(doc, "Son kullanıcılar sisteme iki farklı donanım")
if idx >= 0:
    replace_para_text(doc.paragraphs[idx], 'İstemci Tarafı: Son kullanıcılar sisteme modern web tarayıcıları üzerinden erişmektedir. Uygulama, duyarlı (responsive) tasarım prensipleriyle geliştirilmiş olup masaüstü, tablet ve mobil cihaz tarayıcılarında sorunsuz çalışmaktadır.')
    print(f"[{idx}] 2.2 İstemci tarafı güncellendi")

# Mobil cihazlar satırını güncelle
idx = find_para_by_text(doc, "Mobil Cihazlar: Android ve IOS")
if idx >= 0:
    replace_para_text(doc.paragraphs[idx], '')
    print(f"[{idx}] Mobil cihazlar satırı temizlendi")

idx = find_para_by_text(doc, "Web İstemcileri: Modern web tarayıcılarını")
if idx >= 0:
    replace_para_text(doc.paragraphs[idx], '')
    print(f"[{idx}] Web istemcileri satırı temizlendi (yukarıda birleştirildi)")

# ============================================================
# Kaydet
# ============================================================
doc.save(DST)
print(f"\n✅ Part 1 tamamlandı! Dosya kaydedildi: {DST}")
