"""Part 3: Bolum 3, 4, 5 guncellemeleri."""
import os
from docx import Document

DST = os.path.join(os.path.dirname(__file__), "yeni_updated.docx")
doc = Document(DST)

def rp(para, new_text):
    if not para.runs:
        para.text = new_text
        return
    for run in para.runs:
        run.text = ""
    para.runs[0].text = new_text

def fp(doc, search, start=0):
    for i in range(start, len(doc.paragraphs)):
        if search in doc.paragraphs[i].text:
            return i
    return -1

# === 3.1 Gorev Tamamlama (382-396) ===
i = fp(doc, "Oyunlaştırma tabanlı sistemlerde kullanıcı motivasyonunun")
if i >= 0:
    rp(doc.paragraphs[i], 'Oyunlaştırma tabanlı sistemlerde kullanıcı motivasyonunun sürdürülebilirliği, eylemlerin sonucunda alınan anlık geri bildirimlerle doğrudan ilişkilidir. Bu nedenle sistemde, kullanıcının bir görevi tamamlaması sonrasında elde ettiği ödülün hesaplanması ve ilgili ilerlemenin güncellenmesi için çok katmanlı bir puanlama mekanizması geliştirilmiştir.')
    print(f"[{i}] 3.1 giris")

i = fp(doc, "Tasarlanan mekanizmanın teknik işleyişi")
if i >= 0:
    rp(doc.paragraphs[i], 'Tasarlanan mekanizmanın teknik işleyişi şu adımlardan oluşmuştur:')
    print(f"[{i}] 3.1 adimlar girisi")

i = fp(doc, "İstemci İsteği (Client Request)")
if i >= 0:
    rp(doc.paragraphs[i], '1. İstemci İsteği (Client Request): Süreç, kullanıcının web arayüzünde bir görevi tamamlamasıyla başlamıştır. İstemci, tamamlanan görevin benzersiz kimliğini ve isteğe bağlı olarak mikro günlük (reflection) metnini API\'ye göndermiştir.')
    print(f"[{i}] 3.1 adim 1")

i = fp(doc, "Veri Doğrulama ve Güncelleme")
if i >= 0:
    rp(doc.paragraphs[i], '2. Veri Doğrulama ve Güncelleme: Sunucu (Backend), gelen isteği karşılayarak UserTask tablosunda ilgili kaydı bulmuştur. Görevin durumu "Completed" olarak güncellenmiş, tamamlanma zamanı (CompletedAt) ve kullanıcının o anki enerji/ruh hali anlık görüntüsü (EnergySnapshot, MoodSnapshot) kaydedilmiştir. Ertelenen görevlerde XP Telafi Oranı (XpCompensationRate = 0.70) uygulanmıştır.')
    print(f"[{i}] 3.1 adim 2")

i = fp(doc, "Kategori Bazlı İlerleme Hesaplaması")
if i >= 0:
    rp(doc.paragraphs[i], '3. Kategori Bazlı İlerleme Hesaplaması: GamificationService, tamamlanan görevin hangi kategoriye ait olduğunu belirleyerek UserCategoryProgress tablosundaki ilgili kaydı güncellemiştir. Kazanılan XP, telafi oranı uygulandıktan sonra kategori ilerlemesine eklenmiştir.')
    print(f"[{i}] 3.1 adim 3")

i = fp(doc, "Seviye Atlama Kontrolü (Level Up Logic)")
if i >= 0:
    rp(doc.paragraphs[i], '4. Seviye Atlama Kontrolü (Level Up Logic): Her puan güncellemesinin ardından sistem, kullanıcının mevcut puanının bir üst seviye eşiğini (XpToNextLevel) aşıp aşmadığını kontrol etmiştir.')
    print(f"[{i}] 3.1 adim 4")

i = fp(doc, "Eğer eşik aşılmışsa")
if i >= 0:
    rp(doc.paragraphs[i], '- Eğer eşik aşılmışsa: Kullanıcının o kategorideki seviyesi bir artırılmış, XpToNextLevel değeri 50 XP kadar yükseltilmiştir. Ayrıca kullanıcının streak bilgisi (CurrentStreak) güncellenmiş ve DailyTaskLimit seviyeye göre dinamik olarak ayarlanmıştır.')
    print(f"[{i}] 3.1 adim 4 alt")

i = fp(doc, "Genel Seviye ve Rozet Tetikleyicisi")
if i >= 0:
    rp(doc.paragraphs[i], '5. Genel Seviye, Rozet ve Streak Güncelleme: Kategori bazlı güncellemeler tamamlandıktan sonra, kullanıcının TotalXP değeri üzerinden GeneralLevel yeniden hesaplanmıştır. Ardından CheckAndAwardBadges metodu ile kullanıcının Level, Streak, TaskCount, CategoryMaster ve FirstTime tipindeki rozetlere hak kazanıp kazanmadığı kontrol edilmiş, kazanılan rozetler için bildirim gönderilmiştir.')
    print(f"[{i}] 3.1 adim 5")

i = fp(doc, "Yanıt (Response): Tüm işlemler başarılı olduğunda")
if i >= 0:
    rp(doc.paragraphs[i], '6. Yanıt (Response): Tüm işlemler başarılı olduğunda API, istemciye güncel puan durumunu, atlanan seviyeyi ve kazanılan rozetleri içeren bir yanıt döndürmüştür. Mikro günlük metni varsa, AI mikro servisine gönderilerek duygu analizi yapılmış ve kullanıcıya kişiselleştirilmiş bir içgörü (insight) sunulmuştur.')
    print(f"[{i}] 3.1 adim 6")

# === 3.2 Gorev Atama (413-430) ===
i = fp(doc, "kullanıcının statik bir listeyi takip etmesi yerine")
if i >= 0:
    rp(doc.paragraphs[i], 'Projenin en özgün bileşenlerinden biri olan "Görev Atama Motoru", kullanıcının statik bir listeyi takip etmesi yerine, kişiselleştirilmiş ve dinamik bir görev deneyimi yaşamasını sağlamıştır. Bu motor, DailyQuestWorker arka plan servisi tarafından her gün gece yarısında (UTC) otomatik olarak tetiklenmiş ve kullanıcının ruh hali, enerji seviyesi ve mevcut ilerleme durumuna göre görev havuzu oluşturmuştur.')
    print(f"[{i}] 3.2 giris")

i = fp(doc, "Zamanlayıcı Tetikleyicisi (Scheduler Trigger)")
if i >= 0:
    rp(doc.paragraphs[i], '1. Zamanlayıcı Tetikleyicisi (DailyQuestWorker): Algoritma, her gece yarısı UTC zamanında arka plan servisi olarak otomatik tetiklenmiştir. Worker, paralel batch işleme ile tüm aktif kullanıcıları gruplar halinde (batch size: 100) işlemiştir. Tatil modundaki (IsOnVacation) kullanıcılar atlanmıştır.')
    print(f"[{i}] 3.2 adim 1")

i = fp(doc, "Kullanıcı Profil Analizi")
if i >= 0:
    rp(doc.paragraphs[i], '2. Kullanıcı Profil ve Ruh Hali Analizi: Sistem, aktif kullanıcıların UserCategoryProgress tablosundaki güncel verilerini ve UserDailyMood tablosundaki ruh hali bilgilerini çekmiştir. Enerji seviyesi, müsait süre ve tercih edilen zaman dilimleri görev filtrelemesinde kullanılmıştır.')
    print(f"[{i}] 3.2 adim 2")

i = fp(doc, "Kural Tabanlı Filtreleme (Rule-Based Filtering)")
if i >= 0:
    rp(doc.paragraphs[i], '3. Kural Tabanlı Filtreleme (Rule-Based Filtering): Her kategori için veritabanındaki görev havuzu (AppTask tablosu) taranmış ve şu kurallar uygulanmıştır:')
    print(f"[{i}] 3.2 adim 3")

i = fp(doc, "Kural: Görevin aktif olması")
if i >= 0:
    rp(doc.paragraphs[i], '- Kural: Görevin aktif olması (IsActive=True) VE Görevin zorluk aralığının kullanıcının seviyesini kapsaması (MinLevel <= UserLevel <= MaxLevel) VE Görevin enerji gereksiniminin kullanıcının mevcut enerji seviyesine uygun olması VE Tahmini sürenin kullanıcının müsait süresine sığması.')
    print(f"[{i}] 3.2 adim 3 kural")

i = fp(doc, "Tekrarlılığı Önleme (Exclusion Logic)")
if i >= 0:
    rp(doc.paragraphs[i], '4. Tekrarlılığı Önleme ve Görev Havuzu: Filtrelenen aday görevler arasından, kullanıcının son 48 saat içinde tamamladığı görevler çıkarılmıştır. Uygun görevler, kullanıcının günlük görev limitine (DailyTaskLimit, seviyeye göre 3-10 arası) göre sınırlandırılmış ve "Suggested" statüsüyle görev havuzuna (daily pool) eklenmiştir. Kullanıcı, havuzdan istediği görevleri kabul ederek aktif görev listesine almıştır.')
    print(f"[{i}] 3.2 adim 4")

i = fp(doc, "Atama ve Veritabanı Kaydı")
if i >= 0:
    rp(doc.paragraphs[i], '5. Streak Güncelleme: Worker, aynı zamanda kullanıcıların streak (ardışık gün) durumlarını kontrol etmiştir. Önceki gün en az bir görev tamamlayan kullanıcıların CurrentStreak değeri korunmuş, tamamlamayanların streak\'i sıfırlanmıştır. Tatil modundaki kullanıcıların streak\'i korunmuştur.')
    print(f"[{i}] 3.2 adim 5")

i = fp(doc, "Bildirim Servisi: Yeni görevler atandığında")
if i >= 0:
    rp(doc.paragraphs[i], '6. Bildirim Servisi: Yeni görevler atandığında, kullanıcıya SignalR hub üzerinden gerçek zamanlı bildirim gönderilmiştir. Ayrıca NudgeWorker servisi, gün içinde görevlerini tamamlamamış kullanıcılara proaktif hatırlatma bildirimleri göndermiştir.')
    print(f"[{i}] 3.2 adim 6")

# === 3.4 Is-Zaman Cizelgesi (447-470) ===
i = fp(doc, "kaynakların verimli kullanımı ve geliştirme")
if i >= 0:
    rp(doc.paragraphs[i], '"Yaşam Pusulası" projesinin tasarım aşamasından nihai ürün haline gelmesi sürecinde, kaynakların verimli kullanımı ve geliştirme sürecinin izlenebilirliği için beş aşamalı bir iş-zaman planı uygulanmıştır.')
    print(f"[{i}] 3.4 giris")

i = fp(doc, "Planlanan geliştirme takvimi")
if i >= 0:
    rp(doc.paragraphs[i], 'Uygulanan geliştirme takvimi ve iş paketlerinin detayları şu şekildedir:')
    print(f"[{i}] 3.4 takvim girisi")

# Bootstrap -> React guncelle
i = fp(doc, "Bootstrap kütüphanesi kullanılarak")
if i >= 0:
    rp(doc.paragraphs[i], '- Web Arayüzü: Vite + React 19 + TailwindCSS kullanılarak duyarlı (responsive) SPA tasarımı kodlanmış ve Backend API ile entegrasyonu gerçekleştirilmiştir.')
    print(f"[{i}] 3.4 faz3 web")

i = fp(doc, "React Native teknolojisi ile Android/iOS")
if i >= 0:
    rp(doc.paragraphs[i], '- Sosyal Platform: Arkadaşlık, mücadele, sosyal akış (post/beğeni/yorum), hedef yönetimi ve bildirim sistemi modülleri geliştirilmiştir.')
    print(f"[{i}] 3.4 faz3 mobil->sosyal")

i = fp(doc, "NLP) tabanlı Chatbot servisinin Infrastructure")
if i >= 0:
    rp(doc.paragraphs[i], '- Python FastAPI tabanlı bağımsız AI mikro servisinin geliştirilmesi ve Docker konteyneri olarak dağıtılması gerçekleştirilmiştir.')
    print(f"[{i}] 3.4 faz4 nlp")

i = fp(doc, "Botun, \"görev öner\"")
if i >= 0:
    rp(doc.paragraphs[i], '- SVM tabanlı niyet sınıflandırma (16+ intent) ve duygu analizi modellerinin eğitilmesi, chatbot\'un bağlam-duyarlı yanıt üretmesi sağlanmıştır.')
    print(f"[{i}] 3.4 faz4 bot")

i = fp(doc, "Tüm modüllerin (Web, Mobil, API, Chatbot)")
if i >= 0:
    rp(doc.paragraphs[i], '- Tüm modüllerin (Web, API, AI Servisi, Chatbot) entegrasyon testlerinin gerçekleştirilmesi.')
    print(f"[{i}] 3.4 faz5 test")

# === 3.5 Risk Yonetimi (472-476) ===
i = fp(doc, "karşılaşılabilecek teknik ve operasyonel riskler analiz edilmiş")
if i >= 0:
    rp(doc.paragraphs[i], 'Yazılım geliştirme süreçlerinde karşılaşılabilecek teknik ve operasyonel riskler analiz edilmiş ve projenin sürekliliğini sağlamak adına önleyici tedbirler geliştirilmiştir.')
    print(f"[{i}] 3.5 giris")

i = fp(doc, "Motivasyon Kaybı: Kullanıcıların statik görevlerden")
if i >= 0:
    rp(doc.paragraphs[i], '- Motivasyon Kaybı: Kullanıcıların statik görevlerden sıkılması riskine karşı, ruh hali duyarlı görev önerisi, streak sistemi, mücadele mekanizması ve proaktif bildirimler (NudgeWorker) ile kullanıcı bağlılığı artırılmıştır.\n- Mikro Servis Kesintisi: AI mikro servisinin geçici olarak erişilemez olması durumunda, sistemin temel işlevlerinin (görev atama, tamamlama) çalışmaya devam etmesi sağlanmıştır. Docker healthcheck mekanizması ile servis durumu izlenmiştir.')
    print(f"[{i}] 3.5 motivasyon + mikro servis riski")

# === 3.6 Maliyet (483-486) ===
i = fp(doc, "Yazılım Maliyetleri: Projenin sunucu tarafında")
if i >= 0:
    rp(doc.paragraphs[i], '- Yazılım Maliyetleri: Projenin sunucu tarafında .NET Core (ücretsiz), veritabanı tarafında PostgreSQL (ücretsiz), istemci tarafında Vite + React + TailwindCSS (ücretsiz) ve yapay zeka tarafında Python FastAPI + scikit-learn (ücretsiz) kullanılmıştır. Geliştirme ortamı Docker Compose ile yönetilmiş, MailHog (e-posta testi) ve pgAdmin (veritabanı yönetimi) araçları kullanılmıştır.')
    print(f"[{i}] 3.6 yazilim maliyetleri")

# === BOLUM 4 - Veri Guvenligi ===
i = fp(doc, "Kimlik İbrazı (Credential Submission)")
if i >= 0:
    rp(doc.paragraphs[i], '1. Kimlik İbrazı (Credential Submission): Kullanıcı, web arayüzü üzerinden e-posta ve şifresini /login uç noktasına göndermiştir. E-posta doğrulaması zorunlu tutulmuş, doğrulanmamış hesaplarla giriş engellenmiştir.')
    print(f"[{i}] 4.1 adim 1")

i = fp(doc, "Token Üretimi (Token Generation)")
if i >= 0:
    rp(doc.paragraphs[i], '3. Token Üretimi (Token Generation): Doğrulama başarılı ise sunucu, kullanıcıya özel dijital imzalı bir Erişim Anahtarı (Access Token - JWT) ve Yenileme Anahtarı (Refresh Token) üretmiştir. Her iki token da HttpOnly, Secure ve SameSite özelliklerine sahip cookie olarak istemciye gönderilmiştir. Access Token 15 dakika, Refresh Token 7 gün geçerli olmuştur.')
    print(f"[{i}] 4.1 adim 3")

i = fp(doc, "Durumsuz İletişim (Stateless Communication)")
if i >= 0:
    rp(doc.paragraphs[i], '4. Durumsuz İletişim (Stateless Communication): İstemci, sonraki her API isteğinde token\'ı otomatik olarak cookie aracılığıyla göndermiştir. Token süresi dolduğunda /refresh-token endpoint\'i ile yeni token alınmış, güvenlik için Rate Limiting uygulanmıştır.')
    print(f"[{i}] 4.1 adim 4")

i = fp(doc, "Hareket Halindeki Veri Güvenliği")
if i >= 0:
    rp(doc.paragraphs[i], '- Hareket Halindeki Veri Güvenliği (Data in Transit): React tabanlı istemci ile .NET Core API arasındaki tüm veri trafiği HTTPS üzerinden şifrelenmiştir. CORS politikası ile izin verilen origin\'ler belirlenmiş, yetkisiz kaynakların API\'ye erişimi engellenmiştir.')
    print(f"[{i}] 4.2 transit")

# === BOLUM 5 - Sonuclar ===
i = fp(doc, "5.1. Çalışmanın Konusu ve Yöntem Özeti")
if i >= 0:
    rp(doc.paragraphs[i], '5.1. Çalışmanın Konusu ve Yöntem Özeti')
    print(f"[{i}] 5.1 baslik")

i = fp(doc, "dijital çağda bireylerin yaşadığı motivasyon kaybı ve hedef belirleme sorunlarına çözüm olarak")
if i >= 0:
    rp(doc.paragraphs[i], 'Bu çalışmanın konusu; dijital çağda bireylerin yaşadığı motivasyon kaybı ve hedef belirleme sorunlarına çözüm olarak, oyunlaştırma ve yapay zeka teknolojilerini birleştiren "Yaşam Pusulası" isimli web platformunun geliştirilmesi olmuştur. Platform; eğitim, sağlık, kariyer ve sosyal sorumluluk kategorilerinde görev atama, XP/seviye sistemi, arkadaşlık ve mücadele mekanizmaları ile bireylerin kişisel gelişimini ölçülebilir hale getirmiştir.')
    print(f"[{i}] 5.1 icerik")

i = fp(doc, "Projenin gerçekleştirilmesinde, yazılımın sürdürülebilirliğini sağlamak amacıyla")
if i >= 0:
    rp(doc.paragraphs[i], 'Projenin gerçekleştirilmesinde, yazılımın sürdürülebilirliğini sağlamak amacıyla Temiz Mimari (Clean Architecture) yöntemi uygulanmıştır. Backend .NET Core ve CQRS deseni ile, frontend Vite + React 19 + TailwindCSS ile, yapay zeka asistanı Python FastAPI mikro servis olarak geliştirilmiş ve Docker Compose ile orkestre edilmiştir.')
    print(f"[{i}] 5.1 yontem")

# 5.2
i = fp(doc, "Gerçekleştirilen tasarım ve analiz çalışmaları")
if i >= 0:
    rp(doc.paragraphs[i], 'Gerçekleştirilen geliştirme çalışmaları neticesinde elde edilen temel sonuçlar şunlardır:')
    print(f"[{i}] 5.2 giris")

i = fp(doc, "Bağımsız Mimari Yapısı")
if i >= 0:
    rp(doc.paragraphs[i], '1. Bağımsız Mimari Yapısı: İş kurallarının (Core katmanı) dış dünyadan (Veritabanı, Arayüz) izole edildiği bir mimari yapı başarıyla kurulmuştur. CQRS deseni ve MediatR kütüphanesi ile karmaşık iş akışları yönetilmiş, 6 arka plan servisi ile otonom işlemler gerçekleştirilmiştir.')
    print(f"[{i}] 5.2 sonuc 1")

i = fp(doc, "Kategori Bazlı İlerleme Modeli")
if i >= 0:
    rp(doc.paragraphs[i], '2. Kategori Bazlı İlerleme ve Sosyal Platform: Geleneksel "tekil puan" sistemlerinin aksine; kullanıcının Sağlık, Eğitim ve Sosyal Sorumluluk gibi farklı kategorilerde bağımsız seviye takibi yapan bir model uygulanmıştır. Ek olarak arkadaşlık, mücadele (challenge), sosyal akış (post/beğeni/yorum), hedef yönetimi ve liderlik tablosu gibi sosyal etkileşim mekanizmaları başarıyla geliştirilmiştir.')
    print(f"[{i}] 5.2 sonuc 2")

i = fp(doc, "Dinamik Görev Atama Başarısı")
if i >= 0:
    rp(doc.paragraphs[i], '3. Dinamik Görev Atama ve AI Chatbot: Statik bir görev listesi yerine, kullanıcının o anki seviyesine, ruh haline ve enerji durumuna uygun görevlerin otomatik atanması başarılmıştır. Python FastAPI ile geliştirilen bağımsız AI mikro servisi; SVM tabanlı niyet sınıflandırma (16+ intent), duygu analizi ve kişilik tabanlı yanıt üretimi ile kullanıcıya bağlam-duyarlı destek sunmuştur.')
    print(f"[{i}] 5.2 sonuc 3")

# 5.3
i = fp(doc, "Bireysel gelişimi toplumsal fayda ile bütünleştiren bir platform tasarlama")
if i >= 0:
    rp(doc.paragraphs[i], 'Projenin başlangıcında belirlenen "Bireysel gelişimi toplumsal fayda ile bütünleştiren bir platform geliştirme" hedefine, çalışan bir prototip olarak ulaşılmıştır.')
    print(f"[{i}] 5.3 giris")

i = fp(doc, "Bu çalışma şu aşamada bir Tasarım ve Mimari Öneri")
if i >= 0:
    rp(doc.paragraphs[i], '- Sistemin gerçek kullanıcı verileriyle (Büyük Veri) beslenerek modelin doğruluk ve performans optimizasyonunun yapılması henüz gerçekleştirilmemiş olup, bu çalışma gelecekte pilot uygulama ile değerlendirilecektir.')
    print(f"[{i}] 5.3 eksik 1")

i = fp(doc, "Yapay zeka asistanı (Chatbot), şu an için kural tabanlı senaryolarla")
if i >= 0:
    rp(doc.paragraphs[i], '- Yapay zeka asistanı (Chatbot), SVM tabanlı makine öğrenmesi modelleri ile geliştirilmiş olup; derin öğrenme (Deep Learning) veya büyük dil modelleri (LLM) entegrasyonu gelecek çalışmalar için planlanmıştır.')
    print(f"[{i}] 5.3 eksik 2")

# 5.4 Oneriler
i = fp(doc, "Pilot Uygulama ve Veri Analizi")
if i >= 0:
    rp(doc.paragraphs[i], '1. Pilot Uygulama ve Veri Analizi: Geliştirilen sistemin üniversite öğrencileri üzerinde pilot uygulama olarak test edilmesi ve elde edilen verilerle öneri algoritmalarının iyileştirilmesi önerilmektedir.')
    print(f"[{i}] 5.4 oneri 1")

i = fp(doc, "Yapay Zeka Derinliği: Kural tabanlı öneri sisteminin")
if i >= 0:
    rp(doc.paragraphs[i], '2. Yapay Zeka Derinliği: Mevcut SVM tabanlı modellerin, Derin Öğrenme (Deep Learning) veya Transformer tabanlı büyük dil modelleri (LLM) ile güçlendirilmesi ve kişiselleştirilmiş öneri doğruluğunun artırılması önerilmektedir.')
    print(f"[{i}] 5.4 oneri 2")

i = fp(doc, "Kurumsal Yaygınlaştırma")
if i >= 0:
    rp(doc.paragraphs[i], '3. Kurumsal Yaygınlaştırma ve Mobil Uygulama: Projenin sadece bireysel kullanıcıya değil; eğitim kurumlarına veya şirketlere entegre edilebilecek bir SaaS modeline dönüştürülmesi önerilmektedir. Ayrıca mevcut responsive web uygulamasının React Native veya Flutter ile native mobil uygulamaya dönüştürülmesi, kullanıcı erişimini artıracaktır.')
    print(f"[{i}] 5.4 oneri 3")

# === KAYDET ===
doc.save(DST)
print(f"\nPart 3 tamamlandi! Dosya: {DST}")
