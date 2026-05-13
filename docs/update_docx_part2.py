"""
Part 2: BÖLÜM 2.3.1 (Sistem Mimarisi), 2.3.2 (Veritabanı), 2.3.3 (API Uçları) güncellemeleri.
yeni_updated.docx üzerine çalışır.
"""
import os
from docx import Document

DST = os.path.join(os.path.dirname(__file__), "yeni_updated.docx")
doc = Document(DST)

def replace_para_text(para, new_text):
    if not para.runs:
        para.text = new_text
        return
    for run in para.runs:
        run.text = ""
    para.runs[0].text = new_text

def find_para_by_text(doc, search_text, start_idx=0):
    for i in range(start_idx, len(doc.paragraphs)):
        if search_text in doc.paragraphs[i].text:
            return i
    return -1

# ============================================================
# 2.3.1 Sistem Mimarisi - Katman açıklamaları (paragraf 260-280)
# ============================================================
idx = find_para_by_text(doc, "sistem dört ana katman üzerinde kurgulanmıştır")
if idx >= 0:
    replace_para_text(doc.paragraphs[idx], 'Şekil 2.1\'de görüldüğü üzere sistem beş ana bileşen üzerinde kurgulanmıştır:')
    print(f"[{idx}] 2.3.1 katman girisi guncellendi")

# Core katman
idx = find_para_by_text(doc, "Core (Çekirdek) Katmanı: Projenin merkezi")
if idx >= 0:
    replace_para_text(doc.paragraphs[idx], '1. Core (Çekirdek) Katmanı: Projenin merkezi olan bu katman, diğer tüm katmanlardan bağımsızdır. İki alt bileşenden oluşmuştur:')
    print(f"[{idx}] Core katman guncellendi")

# Domain
idx = find_para_by_text(doc, "Domain: Veritabanı tablolarına karşılık gelen varlıkların")
if idx >= 0:
    replace_para_text(doc.paragraphs[idx], '- Domain: Veritabanı tablolarına karşılık gelen varlıkların (AppUser, AppTask, UserTask, Challenge, Post, UserGoal, Friendship, Notification vb.) ve sistemin değişmez iş kurallarının tanımlandığı katmandır. 26\'dan fazla varlık (entity) bu katmanda tanımlanmıştır.')
    print(f"[{idx}] Domain guncellendi")

# Application
idx = find_para_by_text(doc, "Application: Sistemin senaryolarını")
if idx >= 0:
    replace_para_text(doc.paragraphs[idx], '- Application: Sistemin senaryolarını (Use Cases) içeren iş mantığı katmanıdır. CQRS deseni ve MediatR kütüphanesi ile komut (Command) ve sorgu (Query) işlemleri ayrıştırılmıştır. FluentValidation pipeline behavior ile veri doğrulaması yapılmıştır. Soyut servis arayüzleri (IGamificationService, IAppNotificationService, IUserBehaviorAnalysisService vb.) bu katmanda tanımlanmıştır.')
    print(f"[{idx}] Application guncellendi")

# Infrastructure
idx = find_para_by_text(doc, "Infrastructure (Altyapı) Katmanı")
if idx >= 0:
    replace_para_text(doc.paragraphs[idx], '2. Infrastructure (Altyapı) Katmanı: Çekirdek katmanda tanımlanan soyut arayüzlerin somutlaştırıldığı katmandır. GamificationService (XP/Seviye hesaplama), BehaviorAnalysisService (kullanıcı davranış analizi), EmailService (e-posta gönderimi), TokenService (JWT üretimi) ve SocialService (sosyal etkileşimler) gibi servisler bu katmanda kodlanmıştır.')
    print(f"[{idx}] Infrastructure guncellendi")

# Persistence
idx = find_para_by_text(doc, "Persistence (Veri Erişim) Katmanı")
if idx >= 0:
    replace_para_text(doc.paragraphs[idx], '3. Persistence (Veri Erişim) Katmanı: Veri kalıcılığının ve veritabanı işlemlerinin yönetildiği katmandır. Entity Framework Core ORM aracı ile PostgreSQL veritabanı yönetilmiştir. Generic Repository Pattern (Read/Write ayrımı) ve Code-First yaklaşımı ile veritabanı şeması tasarlanmıştır.')
    print(f"[{idx}] Persistence guncellendi")

# Presentation
idx = find_para_by_text(doc, "Presentation (Sunum) Katmanı")
if idx >= 0:
    replace_para_text(doc.paragraphs[idx], '4. Presentation (Sunum) Katmanı: İstemcilerin sisteme erişim sağladığı katmandır. ASP.NET Core Web API teknolojisi ile RESTful uç noktalar (endpoint) sunulmuştur. Bu katmanda ayrıca altı adet arka plan servisi (BackgroundService) çalışmaktadır: DailyQuestWorker (günlük görev dağıtımı ve streak güncelleme), NudgeWorker (proaktif bildirim), WeeklyReportWorker (haftalık rapor üretimi), GoalReminderWorker (hedef hatırlatma), GoalRevisionWorker (hedef revizyon önerisi) ve ChallengeExpiryWorker (mücadele süresi kontrolü). SignalR hub ile gerçek zamanlı bildirim altyapısı kurulmuştur.')
    print(f"[{idx}] Presentation guncellendi")

# Web Arayuzu — Bootstrap -> React
idx = find_para_by_text(doc, "Web Arayüzü: Kullanıcıların masaüstü")
if idx >= 0:
    replace_para_text(doc.paragraphs[idx], '- Web Arayüzü: Kullanıcıların masaüstü, tablet ve mobil cihaz tarayıcılarından sisteme sorunsuz erişebilmesi için Vite + React 19 + TailwindCSS teknolojileri ile tek sayfa uygulama (SPA) geliştirilmiştir. Bileşen tabanlı mimari sayesinde Dashboard, Görev Listesi, Görev Havuzu, Sosyal Akış, Profil, Hedefler, Raporlar, Rozet Galerisi, Liderlik Tablosu ve Admin Paneli gibi sayfalar oluşturulmuştur. React Router ile istemci taraflı yönlendirme, Axios ile API iletişimi, Recharts ile veri görselleştirme, Formik + Yup ile form yönetimi ve SignalR ile gerçek zamanlı bildirimler sağlanmıştır.')
    print(f"[{idx}] Web arayuzu guncellendi")

# Mobil Uygulama — KALDIR
idx = find_para_by_text(doc, "Mobil Uygulama: Android ve iOS işletim")
if idx >= 0:
    replace_para_text(doc.paragraphs[idx], '5. AI Mikro Servisi: Python FastAPI framework ile geliştirilmiş bağımsız bir yapay zeka servisidir. Docker konteyneri içerisinde çalışmış ve .NET Backend ile HTTP üzerinden haberleşmiştir. Servis; SVM tabanlı niyet sınıflandırma (16+ intent), duygu analizi (positive/negative/neutral), aciliyet tespiti (urgency detection), eylem çıkarımı (action extraction) ve kişilik tabanlı yanıt üretimi (friendly/motivator/coach) modüllerini içermektedir. Ayrıca mikro günlük (reflection) analizi endpoint\'i ile görev tamamlama sonrası kullanıcı düşüncelerinin duygu analizi yapılmıştır.')
    print(f"[{idx}] Mobil uygulama -> AI Mikro Servis guncellendi")

# ============================================================
# 2.3.2 Veritabanı Tasarımı (paragraf 297, 307-317)
# ============================================================
idx = find_para_by_text(doc, "Sistemin veri tutarlılığını sağlamak amacıyla")
if idx >= 0:
    replace_para_text(doc.paragraphs[idx], 'Sistemin veri tutarlılığını sağlamak amacıyla PostgreSQL veritabanı yönetim sistemi üzerinde ilişkisel bir yapı (Relational Database) tasarlanmıştır. Entity Framework Core Code-First yaklaşımı ile oluşturulan veritabanı şeması 26\'dan fazla tabloyu kapsamıştır. Aşağıda temel varlık ilişkileri (ER) diyagramı sunulmuştur.')
    print(f"[{idx}] 2.3.2 giris guncellendi")

# Tablo açıklamaları güncelle
idx = find_para_by_text(doc, "temel tabloların işlevleri şunlardır")
if idx >= 0:
    replace_para_text(doc.paragraphs[idx], 'Şekil 2.2\'deki tasarımda yer alan temel tabloların ve yeni eklenen varlıkların işlevleri şunlardır:')
    print(f"[{idx}] Tablo aciklamalari girisi guncellendi")

# User tablosu
idx = find_para_by_text(doc, "User: Platformun ana aktörü")
if idx >= 0:
    replace_para_text(doc.paragraphs[idx], 'AppUser: Platformun ana aktörü olan kullanıcıların verilerini yönetmek için tasarlanmıştır. ASP.NET Identity üzerine inşa edilmiş bu varlık; güvenli kimlik doğrulama bilgilerinin (şifre hash, e-posta doğrulama, refresh token) yanı sıra oyunlaştırma verilerini (TotalXP, GeneralLevel, CurrentStreak, LongestStreak), davranış istatistiklerini (TaskCompletionRate, AverageSessionMinutes, PreferredTimeSlot), sosyal profil bilgilerini (AvatarId, Bio, Title), tatil/mola modu ayarlarını (IsOnVacation, VacationStartDate) ve onboarding durumunu (IsOnboarded, OnboardingResponse) saklamıştır.')
    print(f"[{idx}] User tablosu guncellendi")

# Task & UserTask
idx = find_para_by_text(doc, "Task & UserTask: Görev yönetim mekanizmasının")
if idx >= 0:
    replace_para_text(doc.paragraphs[idx], 'AppTask ve UserTask: Görev yönetim mekanizmasının merkezini oluşturan bu ikiliden AppTask tablosu; görev tanımlarını (Title, Description), zorluk seviyesini (Difficulty), taban XP değerini (BaseXP), tahmini süreyi (EstimatedMinutes), önerilen zaman dilimlerini (RecommendedTimeSlots - JSONB), enerji seviyesi gereksinimlerini (MinEnergyLevel, MaxEnergyLevel), AI motivasyon metnini (WhyThisTask) ve rozet ipucunu (BadgeHint) içermiştir. UserTask tablosu ise kullanıcı-görev ilişkisini yönetmiş olup; erteleme mantığı (PostponedCount, OriginalDueDate, XpCompensationRate), enerji/ruh hali anlık görüntüsü (MoodSnapshot, EnergySnapshot), mikro günlük (ReflectionText), zamanlayıcı verileri (StartedAt, PausedAt, TotalPausedSeconds) ve sosyal paylaşım alanlarını (ShareToken) barındırmıştır.')
    print(f"[{idx}] Task tablosu guncellendi")

# Category
idx = find_para_by_text(doc, "Category & UserCategoryProgress")
if idx >= 0:
    replace_para_text(doc.paragraphs[idx], 'Category ve UserCategoryProgress: Kullanıcının çok yönlü gelişimini takip etmek amacıyla kurgulanmıştır. Category tablosu gelişim alanlarını (Eğitim, Sağlık, Kariyer, Sosyal Sorumluluk vb.) tanımlamış; UserCategoryProgress ise kullanıcının her kategorideki seviye, mevcut XP ve bir üst seviyeye geçiş eşiği (XpToNextLevel) bilgilerini izlemiştir.')
    print(f"[{idx}] Category guncellendi")

# Badge
idx = find_para_by_text(doc, "Badge & UserBadge: Oyunlaştırma kurgusunu desteklemek")
if idx >= 0:
    replace_para_text(doc.paragraphs[idx], 'Badge ve UserBadge: Oyunlaştırma kurgusunu desteklemek üzere tasarlanan Badge tablosunda; rozet tipi (Level, Streak, TaskCount, CategoryMaster, FirstTime), gerekli değer (RequiredValue), gizli rozet (IsSecret), renk (BadgeColor) ve nadirlik derecesi (Rarity) alanları tanımlanmıştır. UserBadge tablosu ise kullanıcının kazandığı rozetleri ve kazanım tarihlerini kaydetmiştir.')
    print(f"[{idx}] Badge guncellendi")

# ChatBotMessage - güncelle
idx = find_para_by_text(doc, "ChatBotMessage: Yapay zeka asistanı ile kullanıcı")
if idx >= 0:
    replace_para_text(doc.paragraphs[idx], 'ChatBotMessage: Yapay zeka asistanı ile kullanıcı arasındaki etkileşimin sürekliliğini sağlamak amacıyla oluşturulmuştur. Mesaj geçmişi, niyet (intent) ve duygu (sentiment) etiketleri ile birlikte saklanmıştır.\n\nYeni Eklenen Varlıklar: Projenin gelişim sürecinde aşağıdaki varlıklar sisteme eklenmiştir: Challenge ve ChallengeParticipant (kullanıcılar arası mücadele sistemi - TaskCount, XPRace, CategoryFocus tipleri), Friendship (arkadaşlık ilişkileri - Pending, Accepted, Rejected, Blocked durumları), Post, PostTag, PostLike ve PostComment (sosyal akış sistemi - gönderi paylaşma, beğeni, yorum), UserGoal ve GoalMilestone (kişisel hedef yönetimi - otomatik takip, hatırlatma, arşivleme), UserDailyMood (günlük ruh hali kaydı - enerji seviyesi, müsait süre, uyku kalitesi), Notification (bildirim sistemi), WeeklyReport (haftalık performans raporu - AI destekli analiz), Achievement ve UserAchievement (başarım sistemi), OnboardingResponse (kullanıcı karşılama anketi cevapları), UserTaskInteraction (görev etkileşim kaydı) ve Kudos (sosyal beğeni/takdir).')
    print(f"[{idx}] ChatBotMessage + yeni varliklar guncellendi")

# ============================================================
# 2.3.3 API Uçları - Mevcut controller'ları güncelle, yenilerini ekle
# ============================================================

# AuthController güncelle (319-325)
idx = find_para_by_text(doc, "2.3.3.1. AuthController")
if idx >= 0:
    replace_para_text(doc.paragraphs[idx], '2.3.3.1. AuthController')
    print(f"[{idx}] AuthController basligi korundu")

idx = find_para_by_text(doc, "kimlik doğrulama ve yetkilendirme işlemleri")
if idx >= 0:
    replace_para_text(doc.paragraphs[idx], 'Projenin dış dünyaya açılan kapısı olan Sunum (Presentation) katmanında, tüm kimlik doğrulama ve yetkilendirme işlemlerini merkezi olarak yönetmek amacıyla AuthController geliştirilmiştir.')
    print(f"[{idx}] AuthController giris guncellendi")

idx = find_para_by_text(doc, "Kullanıcı kayıt süreçlerinin yönetilmesi adına")
if idx >= 0:
    replace_para_text(doc.paragraphs[idx], 'Kullanıcı kayıt süreci /register uç noktası üzerinden yönetilmiştir. Giriş işlemi /login uç noktası ile gerçekleştirilmiş olup, başarılı kimlik doğrulaması sonrasında JWT token HttpOnly cookie olarak istemciye gönderilmiştir. Güvenliği artırmak amacıyla /verify-email (e-posta doğrulama), /forgot-password (şifre sıfırlama talebi), /reset-password (şifre sıfırlama), /refresh-token (token yenileme) ve /revoke (çıkış) uç noktaları eklenmiştir. Auth endpoint\'lerine Rate Limiting uygulanmıştır.')
    print(f"[{idx}] AuthController register/login guncellendi")

idx = find_para_by_text(doc, "Mevcut kullanıcıların sisteme güvenli erişim")
if idx >= 0:
    replace_para_text(doc.paragraphs[idx], '')
    print(f"[{idx}] AuthController eski login aciklamasi temizlendi")

# TasksController güncelle
idx = find_para_by_text(doc, "2.3.3.5. TasksController")
if idx >= 0:
    replace_para_text(doc.paragraphs[idx], '2.3.3.5. TasksController ve UserTasksController')
    print(f"[{idx}] Tasks basligi guncellendi")

idx = find_para_by_text(doc, "Projenin temel işlevi olan kişisel gelişim görevlerinin")
if idx >= 0:
    replace_para_text(doc.paragraphs[idx], 'Projenin temel işlevi olan kişisel gelişim görevlerinin yaşam döngüsünü yönetmek amacıyla iki ayrı controller geliştirilmiştir: TasksController (görev tanımları CRUD) ve UserTasksController (kullanıcı-görev etkileşimleri).')
    print(f"[{idx}] Tasks giris guncellendi")

idx = find_para_by_text(doc, "Sisteme yeni görev içeriklerinin eklenmesi sürecinin")
if idx >= 0:
    replace_para_text(doc.paragraphs[idx], 'UserTasksController, projenin en kapsamlı controller\'ı olarak geliştirilmiştir. Bu controller; günlük görev havuzu (/daily-pool), görev kabul etme (/accept), görev tamamlama (/complete - mikro günlük desteği ile), görev erteleme (/postpone - XP ceza mekanizmalı), görev başlatma/duraklatma/devam ettirme (/start, /pause, /resume), görevden vazgeçme (/abandon) ve görev paylaşma (/share) uç noktalarını içermiştir.')
    print(f"[{idx}] Tasks detay guncellendi")

idx = find_para_by_text(doc, "Son kullanıcıların sistemdeki aktif görevleri")
if idx >= 0:
    replace_para_text(doc.paragraphs[idx], '')
    print(f"[{idx}] Tasks eski aciklama temizlendi")

# UsersController güncelle
idx = find_para_by_text(doc, "2.3.3.6. UsersController")
if idx >= 0:
    replace_para_text(doc.paragraphs[idx], '2.3.3.6. UsersController ve Yeni Controller\'lar')
    print(f"[{idx}] Users basligi guncellendi")

idx = find_para_by_text(doc, "Projenin sosyal etkileşim ve kişisel takip modüllerini")
if idx >= 0:
    replace_para_text(doc.paragraphs[idx], 'Projenin sosyal etkileşim ve kişisel takip modüllerini yönetmek amacıyla UsersController geliştirilmiştir. Profil görüntüleme (/profile), liderlik tablosu (/leaderboard) ve kullanıcı güncelleme endpoint\'lerini barındırmıştır.')
    print(f"[{idx}] Users giris guncellendi")

idx = find_para_by_text(doc, "Kullanıcıların kendi gelişim süreçlerini")
if idx >= 0:
    replace_para_text(doc.paragraphs[idx], 'Projenin gelişim sürecinde aşağıdaki yeni controller\'lar eklenmiştir: SocialController (arkadaşlık istekleri, engelleme, mücadele daveti/kabul/red, görev paylaşımı, kudos), PostsController (sosyal akış - gönderi oluşturma, beğeni, yorum, silme), GoalsController (hedef oluşturma, ilerleme güncelleme, arşivleme, dashboard), MoodController (günlük ruh hali kaydı, geçmiş sorgulama), ReportsController (haftalık rapor, günlük kırılım), ChatController (AI chatbot iletişimi), NotificationsController (bildirim listeleme, okundu işaretleme), OnboardingController (kullanıcı karşılama anketi) ve AdminController (yönetici işlemleri, analitik).')
    print(f"[{idx}] Yeni controllerlar eklendi")

idx = find_para_by_text(doc, "Oyunlaştırma dinamiklerini güçlendirmek ve kullanıcılar arasında")
if idx >= 0:
    replace_para_text(doc.paragraphs[idx], '')
    print(f"[{idx}] Eski leaderboard aciklamasi temizlendi")

# ============================================================
doc.save(DST)
print(f"\n Part 2 tamamlandi! Dosya kaydedildi: {DST}")
